from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.shortcuts import render, redirect, get_object_or_404
from django.db import models
from .forms import StudentForm
from .models import Student, SchoolClass, Combination, Staff, Department, Subject, Stream, Guardian, OLevelCategory, StudentGuardianRelation
from .forms import (
    StudentForm, StaffUserForm, StaffProfileForm, DepartmentForm, CombinationForm,
    SubjectForm, SchoolClassForm, StreamForm, GuardianForm, OLevelCategoryForm,
)
from apps.accounts.models import User
from .forms import AcademicYearForm, SemesterForm, StudentImportForm
from .models import AcademicYear, Semester
import openpyxl
from datetime import datetime
from .forms import TermForm
from .models import Term
from django.http import HttpResponse
import json
from django.utils import timezone

@login_required
def student_list(request):
    students = Student.objects.exclude(status='GRADUATED').select_related(
        'stream__school_class'
    ).order_by('stream__school_class__order', 'stream__name', 'last_name', 'first_name')
    return render(request, 'sis/student_list.html', {'students': students})


from apps.sis.models import SchoolClass, Combination


@login_required
def student_create(request):
    if request.method == 'POST':
        form = StudentForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'Student registered successfully.')
            return redirect('sis:student_list')
    else:
        form = StudentForm()

    school_classes = SchoolClass.objects.all().order_by('order')
    combinations = Combination.objects.all()
    return render(request, 'sis/student_form.html', {
        'form': form,
        'school_classes': school_classes,
        'combinations': combinations,
    })
from apps.sis.models import SchoolClass, Combination, Student


@login_required
def students_by_class(request):
    classes = SchoolClass.objects.annotate(
        student_count=models.Count(
            'streams__students',
            filter=models.Q(streams__students__status__in=['ACTIVE', 'SUSPENDED', 'TRANSFERRED'])
        )
    ).order_by('order')
    return render(request, 'sis/students_by_class.html', {'classes': classes})

@login_required
def class_students(request, class_id):
    school_class = get_object_or_404(SchoolClass, pk=class_id)
    students = Student.objects.filter(
        stream__school_class=school_class
    ).exclude(status='GRADUATED').select_related('stream')
    return render(request, 'sis/class_students.html', {
        'school_class': school_class,
        'students': students,
    })
@login_required
def student_update(request, pk):
    student = get_object_or_404(Student, pk=pk)
    if request.method == 'POST':
        form = StudentForm(request.POST, instance=student)
        if form.is_valid():
            form.save()
            messages.success(request, 'Student updated successfully.')
            return redirect('sis:student_list')
    else:
        form = StudentForm(instance=student)

    school_classes = SchoolClass.objects.all().order_by('order')
    combinations = Combination.objects.all()
    return render(request, 'sis/student_form.html', {
        'form': form,
        'school_classes': school_classes,
        'combinations': combinations,
        'student': student,
    })


@login_required
def student_delete(request, pk):
    student = get_object_or_404(Student, pk=pk)
    if request.method == 'POST':
        student.delete()
        messages.success(request, 'Student deleted successfully.')
        return redirect('sis:student_list')
    return render(request, 'sis/student_confirm_delete.html', {'student': student})
@login_required
def staff_create(request):
    if request.method == 'POST':
        user_form = StaffUserForm(request.POST)
        profile_form = StaffProfileForm(request.POST, request.FILES)
        if user_form.is_valid() and profile_form.is_valid():
            user = user_form.save()
            staff = profile_form.save(commit=False)
            staff.user = user
            staff.save()
            messages.success(request, f'Staff member {user.get_full_name()} registered successfully.')
            return redirect('sis:staff_list')
    else:
        user_form = StaffUserForm()
        profile_form = StaffProfileForm()

    return render(request, 'sis/staff_form.html', {
        'user_form': user_form,
        'profile_form': profile_form,
    })


@login_required
def staff_list(request):
    staff_members = Staff.objects.select_related('user', 'department').all()
    return render(request, 'sis/staff_list.html', {'staff_members': staff_members})
@login_required
def staff_update(request, pk):
    staff = get_object_or_404(Staff, pk=pk)
    if request.method == 'POST':
        profile_form = StaffProfileForm(request.POST, request.FILES, instance=staff)
        if profile_form.is_valid():
            profile_form.save()
            messages.success(request, 'Staff member updated successfully.')
            return redirect('sis:staff_list')
    else:
        profile_form = StaffProfileForm(instance=staff)

    return render(request, 'sis/staff_update_form.html', {
        'profile_form': profile_form,
        'staff': staff,
    })


@login_required
def staff_delete(request, pk):
    staff = get_object_or_404(Staff, pk=pk)
    if request.method == 'POST':
        user = staff.user
        staff.delete()
        if user:
            user.delete()
        messages.success(request, 'Staff member deleted successfully.')
        return redirect('sis:staff_list')
    return render(request, 'sis/staff_confirm_delete.html', {'staff': staff})
@login_required
def department_list(request):
    departments = Department.objects.all()
    return render(request, 'sis/department_list.html', {'departments': departments})


@login_required
def department_create(request):
    if request.method == 'POST':
        form = DepartmentForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'Department created successfully.')
            return redirect('sis:department_list')
    else:
        form = DepartmentForm()
    return render(request, 'sis/department_form.html', {'form': form})


@login_required
def department_update(request, pk):
    department = get_object_or_404(Department, pk=pk)
    if request.method == 'POST':
        form = DepartmentForm(request.POST, instance=department)
        if form.is_valid():
            form.save()
            messages.success(request, 'Department updated successfully.')
            return redirect('sis:department_list')
    else:
        form = DepartmentForm(instance=department)
    return render(request, 'sis/department_form.html', {'form': form, 'department': department})


@login_required
def department_delete(request, pk):
    department = get_object_or_404(Department, pk=pk)
    if request.method == 'POST':
        department.delete()
        messages.success(request, 'Department deleted successfully.')
        return redirect('sis:department_list')
    return render(request, 'sis/department_confirm_delete.html', {'department': department})
@login_required
def combination_list(request):
    combinations = Combination.objects.all()
    return render(request, 'sis/combination_list.html', {'combinations': combinations})

@login_required
def combination_create(request):
    if request.method == 'POST':
        form = CombinationForm(request.POST)
        if form.is_valid():
            combo = form.save(commit=False)
            combo.added_by = request.user
            combo.save()
            messages.success(request, 'Combination created successfully.')
            return redirect('sis:combination_list')
    else:
        form = CombinationForm()
    return render(request, 'sis/combination_form.html', {'form': form})


@login_required
def combination_update(request, pk):
    combination = get_object_or_404(Combination, pk=pk)
    if request.method == 'POST':
        form = CombinationForm(request.POST, instance=combination)
        if form.is_valid():
            form.save()
            messages.success(request, 'Combination updated successfully.')
            return redirect('sis:combination_list')
    else:
        form = CombinationForm(instance=combination)
    return render(request, 'sis/combination_form.html', {'form': form, 'combination': combination})


@login_required
def combination_delete(request, pk):
    combination = get_object_or_404(Combination, pk=pk)
    if request.method == 'POST':
        combination.delete()
        messages.success(request, 'Combination deleted successfully.')
        return redirect('sis:combination_list')
    return render(request, 'sis/combination_confirm_delete.html', {'combination': combination})

@login_required
def subject_list(request):
    subjects = Subject.objects.select_related('department').all()
    return render(request, 'sis/subject_list.html', {'subjects': subjects})


@login_required
def subject_create(request):
    if request.method == 'POST':
        form = SubjectForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'Subject created successfully.')
            return redirect('sis:subject_list')
    else:
        form = SubjectForm()
    return render(request, 'sis/subject_form.html', {'form': form})


@login_required
def subject_update(request, pk):
    subject = get_object_or_404(Subject, pk=pk)
    if request.method == 'POST':
        form = SubjectForm(request.POST, instance=subject)
        if form.is_valid():
            form.save()
            messages.success(request, 'Subject updated successfully.')
            return redirect('sis:subject_list')
    else:
        form = SubjectForm(instance=subject)
    return render(request, 'sis/subject_form.html', {'form': form, 'subject': subject})


@login_required
def subject_delete(request, pk):
    subject = get_object_or_404(Subject, pk=pk)
    if request.method == 'POST':
        subject.delete()
        messages.success(request, 'Subject deleted successfully.')
        return redirect('sis:subject_list')
    return render(request, 'sis/subject_confirm_delete.html', {'subject': subject})
@login_required
def schoolclass_list(request):
    classes = SchoolClass.objects.all().order_by('order')
    return render(request, 'sis/schoolclass_list.html', {'classes': classes})


@login_required
def schoolclass_create(request):
    if request.method == 'POST':
        form = SchoolClassForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'Class created successfully.')
            return redirect('sis:schoolclass_list')
    else:
        form = SchoolClassForm()
    return render(request, 'sis/schoolclass_form.html', {'form': form})


@login_required
def schoolclass_update(request, pk):
    school_class = get_object_or_404(SchoolClass, pk=pk)
    if request.method == 'POST':
        form = SchoolClassForm(request.POST, instance=school_class)
        if form.is_valid():
            form.save()
            messages.success(request, 'Class updated successfully.')
            return redirect('sis:schoolclass_list')
    else:
        form = SchoolClassForm(instance=school_class)
    return render(request, 'sis/schoolclass_form.html', {'form': form, 'school_class': school_class})


@login_required
def schoolclass_delete(request, pk):
    school_class = get_object_or_404(SchoolClass, pk=pk)
    if request.method == 'POST':
        school_class.delete()
        messages.success(request, 'Class deleted successfully.')
        return redirect('sis:schoolclass_list')
    return render(request, 'sis/schoolclass_confirm_delete.html', {'school_class': school_class})
@login_required
def stream_list(request):
    streams = Stream.objects.select_related('school_class', 'combination').all()
    return render(request, 'sis/stream_list.html', {'streams': streams})


@login_required
def stream_create(request):
    if request.method == 'POST':
        form = StreamForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'Stream created successfully.')
            return redirect('sis:stream_list')
    else:
        form = StreamForm()
    return render(request, 'sis/stream_form.html', {'form': form})


@login_required
def stream_update(request, pk):
    stream = get_object_or_404(Stream, pk=pk)
    if request.method == 'POST':
        form = StreamForm(request.POST, instance=stream)
        if form.is_valid():
            form.save()
            messages.success(request, 'Stream updated successfully.')
            return redirect('sis:stream_list')
    else:
        form = StreamForm(instance=stream)
    return render(request, 'sis/stream_form.html', {'form': form, 'stream': stream})


@login_required
def stream_delete(request, pk):
    stream = get_object_or_404(Stream, pk=pk)
    if request.method == 'POST':
        stream.delete()
        messages.success(request, 'Stream deleted successfully.')
        return redirect('sis:stream_list')
    return render(request, 'sis/stream_confirm_delete.html', {'stream': stream})
@login_required
def guardian_list(request):
    guardians = Guardian.objects.all()
    return render(request, 'sis/guardian_list.html', {'guardians': guardians})


@login_required
def guardian_create(request):
    if request.method == 'POST':
        form = GuardianForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'Guardian registered successfully.')
            return redirect('sis:guardian_list')
    else:
        form = GuardianForm()
    return render(request, 'sis/guardian_form.html', {'form': form})


@login_required
def guardian_update(request, pk):
    guardian = get_object_or_404(Guardian, pk=pk)
    if request.method == 'POST':
        form = GuardianForm(request.POST, instance=guardian)
        if form.is_valid():
            form.save()
            messages.success(request, 'Guardian updated successfully.')
            return redirect('sis:guardian_list')
    else:
        form = GuardianForm(instance=guardian)
    return render(request, 'sis/guardian_form.html', {'form': form, 'guardian': guardian})


@login_required
def guardian_delete(request, pk):
    guardian = get_object_or_404(Guardian, pk=pk)
    if request.method == 'POST':
        guardian.delete()
        messages.success(request, 'Guardian deleted successfully.')
        return redirect('sis:guardian_list')
    return render(request, 'sis/guardian_confirm_delete.html', {'guardian': guardian})
@login_required
def leavers_years(request):
    years = Student.objects.exclude(
        leaving_year__isnull=True
    ).values_list('leaving_year', flat=True).distinct().order_by('-leaving_year')
    return render(request, 'sis/leavers_years.html', {'years': years})


@login_required
def leavers_year_detail(request, year):
    has_form4 = Student.objects.filter(leaving_year=year, leaving_class__icontains='4').exists()
    has_form6 = Student.objects.filter(leaving_year=year, leaving_class__icontains='6').exists()
    return render(request, 'sis/leavers_year_detail.html', {
        'year': year,
        'has_form4': has_form4,
        'has_form6': has_form6,
    })


@login_required
def leavers_batch(request, year, form):
    students = Student.objects.filter(
        leaving_year=year, leaving_class__icontains=form
    ).select_related('stream')
    return render(request, 'sis/leavers_batch.html', {
        'year': year,
        'form': form,
        'students': students,
    })


@login_required
def olevelcategory_list(request):
    categories = OLevelCategory.objects.all()
    return render(request, 'sis/olevelcategory_list.html', {'categories': categories})


@login_required
def olevelcategory_create(request):
    if request.method == 'POST':
        form = OLevelCategoryForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'O-Level category created successfully.')
            return redirect('sis:olevelcategory_list')
    else:
        form = OLevelCategoryForm()
    return render(request, 'sis/olevelcategory_form.html', {'form': form})


@login_required
def olevelcategory_update(request, pk):
    category = get_object_or_404(OLevelCategory, pk=pk)
    if request.method == 'POST':
        form = OLevelCategoryForm(request.POST, instance=category)
        if form.is_valid():
            form.save()
            messages.success(request, 'O-Level category updated successfully.')
            return redirect('sis:olevelcategory_list')
    else:
        form = OLevelCategoryForm(instance=category)
    return render(request, 'sis/olevelcategory_form.html', {'form': form, 'category': category})


@login_required
def olevelcategory_delete(request, pk):
    category = get_object_or_404(OLevelCategory, pk=pk)
    if request.method == 'POST':
        category.delete()
        messages.success(request, 'O-Level category deleted successfully.')
        return redirect('sis:olevelcategory_list')
    return render(request, 'sis/olevelcategory_confirm_delete.html', {'category': category})

@login_required
def academicyear_list(request):
    years = AcademicYear.objects.all()
    return render(request, 'sis/academicyear_list.html', {'years': years})


@login_required
def academicyear_create(request):
    if request.method == 'POST':
        form = AcademicYearForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'Academic year created successfully.')
            return redirect('sis:academicyear_list')
    else:
        form = AcademicYearForm()
    return render(request, 'sis/academicyear_form.html', {'form': form})


@login_required
def academicyear_update(request, pk):
    year = get_object_or_404(AcademicYear, pk=pk)
    if request.method == 'POST':
        form = AcademicYearForm(request.POST, instance=year)
        if form.is_valid():
            form.save()
            messages.success(request, 'Academic year updated successfully.')
            return redirect('sis:academicyear_list')
    else:
        form = AcademicYearForm(instance=year)
    return render(request, 'sis/academicyear_form.html', {'form': form, 'year': year})


from django.db.models import ProtectedError


@login_required
def academicyear_delete(request, pk):
    year = get_object_or_404(AcademicYear, pk=pk)
    if request.method == 'POST':
        try:
            year.delete()
            messages.success(request, 'Academic year deleted successfully.')
        except ProtectedError:
            messages.error(
                request,
                f'Cannot delete {year.year} — it has Fee Structures with Invoices/Payments already '
                'recorded against it. Financial records are protected from deletion to preserve history.'
            )
        return redirect('sis:academicyear_list')
    return render(request, 'sis/academicyear_confirm_delete.html', {'year': year})


@login_required
def semester_list(request):
    semesters = Semester.objects.select_related('academic_year').all()
    return render(request, 'sis/semester_list.html', {'semesters': semesters})


@login_required
def semester_create(request):
    if request.method == 'POST':
        form = SemesterForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'Semester created successfully.')
            return redirect('sis:semester_list')
    else:
        form = SemesterForm()
    return render(request, 'sis/semester_form.html', {'form': form})


@login_required
def semester_update(request, pk):
    semester = get_object_or_404(Semester, pk=pk)
    if request.method == 'POST':
        form = SemesterForm(request.POST, instance=semester)
        if form.is_valid():
            form.save()
            messages.success(request, 'Semester updated successfully.')
            return redirect('sis:semester_list')
    else:
        form = SemesterForm(instance=semester)
    return render(request, 'sis/semester_form.html', {'form': form, 'semester': semester})


@login_required
def semester_delete(request, pk):
    semester = get_object_or_404(Semester, pk=pk)
    if request.method == 'POST':
        semester.delete()
        messages.success(request, 'Semester deleted successfully.')
        return redirect('sis:semester_list')
    return render(request, 'sis/semester_confirm_delete.html', {'semester': semester})
@login_required
def class_teachers_manage(request):
    classes = SchoolClass.objects.select_related('class_teacher').order_by('order')
    teachers = User.objects.filter(role=User.Role.TEACHER)

    if request.method == 'POST':
        class_id = request.POST.get('class_id')
        teacher_id = request.POST.get('teacher_id') or None
        school_class = get_object_or_404(SchoolClass, pk=class_id)
        school_class.class_teacher_id = teacher_id
        school_class.save()
        messages.success(request, f'Class teacher for {school_class.name} updated successfully.')
        return redirect('sis:class_teachers_manage')

    return render(request, 'sis/class_teachers_manage.html', {'classes': classes, 'teachers': teachers})

@login_required
def student_import_template(request):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Students"

    headers = [
        "First Name", "Middle Name", "Last Name", "Gender (M/F)",
        "Date of Birth (YYYY-MM-DD)", "Class", "Stream",
        "Combination Code (A-Level only)", "Guardian Phone Number",
        "Status (optional, defaults to ACTIVE)",
    ]
    ws.append(headers)

    # Example row shows the expected format without needing extra instructions
    ws.append([
        "John", "A.", "Doe", "M", "2008-05-14",
        "Form 1", "A", "", "0712345678", "ACTIVE",
    ])

    for col_num, header in enumerate(headers, start=1):
        ws.column_dimensions[ws.cell(row=1, column=col_num).column_letter].width = max(len(header) * 1.1, 15)

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="student_import_template.xlsx"'
    wb.save(response)
    return response

@login_required
def student_import(request):
    if request.method == 'POST':
        form = StudentImportForm(request.POST, request.FILES)
        if form.is_valid():
            wb = openpyxl.load_workbook(request.FILES['excel_file'])
            ws = wb.active

            created_count = 0
            skipped_rows = []

            for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                if not row or not row[0]:
                    continue

                (first_name, middle_name, last_name, gender, dob, class_name,
                 stream_name, combination_code, guardian_phone, status_text) = (list(row) + [None] * 10)[:10]

                school_class = SchoolClass.objects.filter(name__iexact=str(class_name).strip()).first() if class_name else None
                if not school_class:
                    skipped_rows.append(f"Row {row_num}: class '{class_name}' not found")
                    continue

                stream = Stream.objects.filter(
                    school_class=school_class, name__iexact=str(stream_name).strip()
                ).first() if stream_name else None
                if not stream:
                    skipped_rows.append(f"Row {row_num}: stream '{stream_name}' not found in {class_name}")
                    continue

                combination = None
                if combination_code:
                    combination = Combination.objects.filter(code__iexact=str(combination_code).strip()).first()
                    if not combination:
                        skipped_rows.append(f"Row {row_num}: combination '{combination_code}' not found")
                        continue

                date_of_birth = None
                if dob:
                    if isinstance(dob, datetime):
                        date_of_birth = dob.date()
                    else:
                        try:
                            date_of_birth = datetime.strptime(str(dob), '%Y-%m-%d').date()
                        except ValueError:
                            pass

                status = 'ACTIVE'
                if status_text and str(status_text).strip().upper() in dict(Student.Status.choices):
                    status = str(status_text).strip().upper()

                student = Student.objects.create(
                    first_name=str(first_name).strip() if first_name else '',
                    middle_name=str(middle_name).strip() if middle_name else '',
                    last_name=str(last_name).strip() if last_name else '',
                    gender=str(gender).strip()[:1].upper() if gender else '',
                    date_of_birth=date_of_birth,
                    stream=stream,
                    combination=combination,
                    status=status,
                )

                if guardian_phone:
                    guardian, _ = Guardian.objects.get_or_create(
                        phone_number=str(guardian_phone).strip(),
                        defaults={'first_name': 'Guardian', 'last_name': f'of {student.first_name}'}
                    )
                    StudentGuardianRelation.objects.get_or_create(
                        student=student, guardian=guardian,
                        defaults={'relationship': 'GUARDIAN', 'is_primary_contact': True}
                    )

                created_count += 1

            if created_count:
                messages.success(request, f'{created_count} student(s) imported successfully.')
            if skipped_rows:
                messages.warning(request, f'{len(skipped_rows)} row(s) skipped: ' + '; '.join(skipped_rows[:5]))

            return redirect('sis:student_list')
    else:
        form = StudentImportForm()

    return render(request, 'sis/student_import.html', {'form': form})
    
@login_required
def term_list(request):
    terms = Term.objects.select_related('semester', 'semester__academic_year').all()
    return render(request, 'sis/term_list.html', {'terms': terms})


@login_required
def term_create(request):
    if request.method == 'POST':
        form = TermForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'Term created successfully.')
            return redirect('sis:term_list')
    else:
        form = TermForm()
    return render(request, 'sis/term_form.html', {'form': form})


@login_required
def term_update(request, pk):
    term = get_object_or_404(Term, pk=pk)
    if request.method == 'POST':
        form = TermForm(request.POST, instance=term)
        if form.is_valid():
            form.save()
            messages.success(request, 'Term updated successfully.')
            return redirect('sis:term_list')
    else:
        form = TermForm(instance=term)
    return render(request, 'sis/term_form.html', {'form': form, 'term': term})


@login_required
def term_delete(request, pk):
    term = get_object_or_404(Term, pk=pk)
    if request.method == 'POST':
        term.delete()
        messages.success(request, 'Term deleted successfully.')
        return redirect('sis:term_list')
    return render(request, 'sis/term_confirm_delete.html', {'term': term})

@login_required
def student_detail(request, pk):
    student = get_object_or_404(
        Student.objects.select_related('stream__school_class', 'combination'), pk=pk
    )
    guardians = student.guardians.all()
    subjects = Subject.objects.filter(
        levels__in=['BOTH', student.stream.school_class.level] if student.stream else []
    )
    return render(request, 'sis/student_detail.html', {
        'student': student,
        'guardians': guardians,
        'subjects': subjects,
    })

@login_required
def staff_detail(request, pk):
    staff = get_object_or_404(
        Staff.objects.select_related('user', 'department'), pk=pk
    )
    subjects_taught = staff.user.subjects_taught.all()
    classes_managed = SchoolClass.objects.filter(class_teacher=staff.user)
    departments_headed = staff.departments_headed.all()
    return render(request, 'sis/staff_detail.html', {
        'staff': staff,
        'subjects_taught': subjects_taught,
        'classes_managed': classes_managed,
        'departments_headed': departments_headed,
    })

from apps.core.decorators import role_required


@login_required
@role_required('TEACHER')
def my_class(request):
    classes = SchoolClass.objects.filter(class_teacher=request.user).order_by('order')

    class_data = []
    for school_class in classes:
        students = Student.objects.filter(
            stream__school_class=school_class
        ).exclude(status='GRADUATED').select_related('stream').order_by('last_name', 'first_name')
        class_data.append({
            'school_class': school_class,
            'students': students,
            'student_count': students.count(),
        })

    staff_profile = getattr(request.user, 'staff_profile', None)

    return render(request, 'sis/my_class.html', {
        'class_data': class_data,
        'staff_profile': staff_profile,
    })

from .forms import TeachingScheduleEntryForm, LessonPlanForm, TopicForm, SubTopicForm
from .models import TeachingScheduleEntry, LessonPlan, Topic, SubTopic


# ---------- Teaching Schedule ----------

@login_required
@role_required('TEACHER')
def teaching_schedule_list(request):
    entries = TeachingScheduleEntry.objects.filter(teacher=request.user).select_related('stream', 'subject')

    if request.method == 'POST':
        form = TeachingScheduleEntryForm(request.POST, teacher=request.user)
        if form.is_valid():
            entry = form.save(commit=False)
            entry.teacher = request.user
            entry.save()
            messages.success(request, 'Schedule entry added.')
            return redirect('sis:teaching_schedule_list')
    else:
        form = TeachingScheduleEntryForm(teacher=request.user)

    days_order = ['MON', 'TUE', 'WED', 'THU', 'FRI', 'SAT']
    schedule_by_day = {day: [] for day in days_order}
    for entry in entries:
        schedule_by_day[entry.day].append(entry)

    return render(request, 'sis/teaching_schedule_list.html', {
        'form': form,
        'schedule_by_day': schedule_by_day,
    })


@login_required
@role_required('TEACHER')
def teaching_schedule_delete(request, pk):
    entry = get_object_or_404(TeachingScheduleEntry, pk=pk, teacher=request.user)
    if request.method == 'POST':
        entry.delete()
        messages.success(request, 'Schedule entry removed.')
        return redirect('sis:teaching_schedule_list')
    return render(request, 'sis/teaching_schedule_confirm_delete.html', {'entry': entry})


# ---------- Lesson Plans ----------

@login_required
@role_required('TEACHER')
def lesson_plan_list(request):
    plans = LessonPlan.objects.filter(teacher=request.user).select_related('stream', 'subject')
    return render(request, 'sis/lesson_plan_list.html', {'plans': plans})


@login_required
@role_required('TEACHER')
def lesson_plan_create(request):
    if request.method == 'POST':
        form = LessonPlanForm(request.POST, teacher=request.user)
        if form.is_valid():
            plan = form.save(commit=False)
            plan.teacher = request.user
            plan.save()
            messages.success(request, 'Lesson plan created. Now add topics and subtopics.')
            return redirect('sis:lesson_plan_detail', pk=plan.pk)
    else:
        form = LessonPlanForm(teacher=request.user)
    return render(request, 'sis/lesson_plan_form.html', {'form': form})


@login_required
@role_required('TEACHER')
def lesson_plan_detail(request, pk):
    plan = get_object_or_404(LessonPlan, pk=pk, teacher=request.user)
    topics = plan.topics.prefetch_related('subtopics')
    topic_form = TopicForm()
    subtopic_form = SubTopicForm()

    chart_labels = [t.title for t in topics]
    chart_progress = [t.progress_percent() for t in topics]

    return render(request, 'sis/lesson_plan_detail.html', {
        'plan': plan,
        'topics': topics,
        'topic_form': topic_form,
        'subtopic_form': subtopic_form,
        'chart_labels_json': json.dumps(chart_labels),
        'chart_progress_json': json.dumps(chart_progress),
    })


@login_required
@role_required('TEACHER')
def lesson_plan_delete(request, pk):
    plan = get_object_or_404(LessonPlan, pk=pk, teacher=request.user)
    if request.method == 'POST':
        plan.delete()
        messages.success(request, 'Lesson plan deleted.')
        return redirect('sis:lesson_plan_list')
    return render(request, 'sis/lesson_plan_confirm_delete.html', {'plan': plan})


@login_required
@role_required('TEACHER')
def topic_add(request, plan_pk):
    plan = get_object_or_404(LessonPlan, pk=plan_pk, teacher=request.user)
    if request.method == 'POST':
        form = TopicForm(request.POST)
        if form.is_valid():
            topic = form.save(commit=False)
            topic.lesson_plan = plan
            topic.save()
            messages.success(request, 'Topic added.')
    return redirect('sis:lesson_plan_detail', pk=plan.pk)


@login_required
@role_required('TEACHER')
def topic_delete(request, pk):
    topic = get_object_or_404(Topic, pk=pk, lesson_plan__teacher=request.user)
    plan_pk = topic.lesson_plan.pk
    if request.method == 'POST':
        topic.delete()
        messages.success(request, 'Topic deleted.')
    return redirect('sis:lesson_plan_detail', pk=plan_pk)


@login_required
@role_required('TEACHER')
def subtopic_add(request, topic_pk):
    topic = get_object_or_404(Topic, pk=topic_pk, lesson_plan__teacher=request.user)
    if request.method == 'POST':
        form = SubTopicForm(request.POST)
        if form.is_valid():
            subtopic = form.save(commit=False)
            subtopic.topic = topic
            subtopic.save()
            messages.success(request, 'Subtopic added.')
    return redirect('sis:lesson_plan_detail', pk=topic.lesson_plan.pk)


@login_required
@role_required('TEACHER')
def subtopic_toggle(request, pk):
    subtopic = get_object_or_404(SubTopic, pk=pk, topic__lesson_plan__teacher=request.user)
    if request.method == 'POST':
        subtopic.is_taught = not subtopic.is_taught
        subtopic.date_taught = timezone.now().date() if subtopic.is_taught else None
        subtopic.save()
    return redirect('sis:lesson_plan_detail', pk=subtopic.topic.lesson_plan.pk)


@login_required
@role_required('TEACHER')
def subtopic_delete(request, pk):
    subtopic = get_object_or_404(SubTopic, pk=pk, topic__lesson_plan__teacher=request.user)
    plan_pk = subtopic.topic.lesson_plan.pk
    if request.method == 'POST':
        subtopic.delete()
        messages.success(request, 'Subtopic deleted.')
    return redirect('sis:lesson_plan_detail', pk=plan_pk)

